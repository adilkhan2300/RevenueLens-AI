"""
Script to generate a comprehensive, executive & academic Word (.docx) report
explaining RevenueLens AI in full detail.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding in a table cell (values in twips, 1/20 of a pt)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m_name, m_val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def create_report():
    doc = Document()

    # Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Palette
    NAVY = RGBColor(30, 58, 138)       # #1E3A8A
    BLUE = RGBColor(37, 99, 235)       # #2563EB
    DARK_GRAY = RGBColor(31, 41, 55)   # #1F2937
    MUTED_GRAY = RGBColor(100, 116, 139) # #64748B

    # Normal Style configuration
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = DARK_GRAY
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # TITLE & COVER BANNER
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("RevenueLens AI")
    run_title.font.name = 'Helvetica'
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    run_sub = p_sub.add_run("Marketing-to-Revenue Attribution & Funnel Intelligence Platform")
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = BLUE

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(24)
    run_meta = p_meta.add_run("Comprehensive Technical, Analytical & Strategic Project Report\n"
                              "Academic Track: MBA Marketing & Data Science / AI Analytics")
    run_meta.font.size = Pt(10)
    run_meta.font.color.rgb = MUTED_GRAY

    # Divider Table
    t_div = doc.add_table(rows=1, cols=1)
    t_div.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_background(t_div.rows[0].cells[0], "2563EB")
    t_div.rows[0].cells[0].paragraphs[0].paragraph_format.space_before = Pt(2)
    t_div.rows[0].cells[0].paragraphs[0].paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function for headings
    def add_custom_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        if level == 1:
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(text)
            r.font.name = 'Helvetica'
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = NAVY
        elif level == 2:
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.font.name = 'Helvetica'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = BLUE
        else:
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = DARK_GRAY
        return p

    def add_callout(text, prefix="KEY STRATEGIC TAKEAWAY: "):
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r_pre = p.add_run(prefix)
        r_pre.font.bold = True
        r_pre.font.color.rgb = BLUE
        r_txt = p.add_run(text)
        r_txt.font.italic = True
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY & PROBLEM STATEMENT
    # -------------------------------------------------------------
    add_custom_heading("1. Executive Summary & Problem Context", 1)

    doc.add_paragraph(
        "Modern digital enterprises and B2B/B2C organizations frequently face a costly operational dilemma known as the "
        "'Marketing-to-Revenue Growth Paradox': despite aggressive increases in marketing expenditure and record volumes of "
        "top-of-funnel inbound leads, top-line recognized revenue growth remains sluggish or stagnant. Marketing departments "
        "celebrate low Cost Per Lead (CPL) and surging form fills, while Sales departments report that inbound prospects are "
        "unqualified, disengaged, or impossible to convert into paying customers."
    )

    doc.add_paragraph(
        "RevenueLens AI was engineered to bridge this critical divide. Designed as a high-caliber analytics platform for CMOs, "
        "Chief Revenue Officers (CROs), and Data Science practitioners, the system provides end-to-end visibility across the entire "
        "commercial pipeline. By tracking customer interactions from initial ad exposure through lead qualification, opportunity stages, "
        "and ultimate cash recognition, RevenueLens AI isolates funnel choke points, measures multi-touch revenue attribution, "
        "detects capital misallocation, and delivers mathematically grounded turnaround recommendations."
    )

    add_callout(
        "Commercial growth fails not due to lack of inbound leads, but due to 'Vanity Metrics' masking poor lead qualification, "
        "speed-to-lead latency, and marketing spend concentrated in channels with negligible deal velocity."
    )

    # -------------------------------------------------------------
    # 2. SYSTEM ARCHITECTURE & TECHNICAL STACK
    # -------------------------------------------------------------
    add_custom_heading("2. System Architecture & Technical Stack", 1)

    doc.add_paragraph(
        "RevenueLens AI is architected with a decoupled modular design pattern in Python, adhering to enterprise software engineering "
        "standards and avoiding monolithic scripting. Each core function (KPI calculation, funnel progression, multi-touch attribution, "
        "machine learning classification, heuristic leakage detection, scenario forecasting, and PDF compilation) is isolated in a "
        "reusable, unit-tested engine module."
    )

    # Tech Stack Table
    t_stack = doc.add_table(rows=1, cols=3)
    t_stack.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Architecture Layer", "Primary Technology", "Strategic Capability & Role"]
    for i, h in enumerate(headers):
        cell = t_stack.rows[0].cells[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    stack_data = [
        ("Frontend & Dashboard", "Streamlit (Modern SaaS layout)", "Interactive multi-page UI, session management, dynamic KPI cards, and state persistence."),
        ("Data Visualization", "Plotly (Go & Express)", "Responsive 9-stage funnel charts, bubble quadrant plots, ROC curves, and grouped comparative bars."),
        ("Data Science & Modeling", "Scikit-learn, Pandas, NumPy", "ColumnTransformer pipelines, Logistic Regression, Random Forest, stratified cross-validation, and metrics."),
        ("Data Persistence", "SQLite3 (Relational DB)", "Local relational schema storing campaigns, leads, touchpoint journeys, and closed sales opportunities."),
        ("Intelligence & Synthesis", "Hybrid Heuristics + OpenAI API", "Deterministic 10-point rule engine (100% offline) with optional LLM narrative generation."),
        ("Executive Reporting", "ReportLab Platypus", "Direct compilation of publication-quality C-suite executive briefing PDF documents."),
        ("Automated Quality Assur.", "Pytest", "Automated test suite verifying attribution normalization, zero-division resilience, and data leakage checks.")
    ]

    for row_idx, (layer, tech, role) in enumerate(stack_data):
        row = t_stack.add_row()
        bg_col = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate([layer, tech, role]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, 100, 100, 140, 140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c_idx == 0:
                p.add_run(val).font.bold = True
            else:
                p.add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 3. DETAILED MODULE-BY-MODULE WALKTHROUGH
    # -------------------------------------------------------------
    add_custom_heading("3. Core Modules & Strategic Capabilities", 1)

    # Module 1
    add_custom_heading("Module 1: Executive Dashboard", 2)
    doc.add_paragraph(
        "The Executive Dashboard acts as the primary cockpit for commercial leadership. It visualizes 15 mission-critical marketing "
        "and revenue KPIs, including Total Spend, Attributed Revenue, Return on Ad Spend (ROAS), Marketing ROI, Average Deal Value (ACV), "
        "Cost Per Lead (CPL), Customer Acquisition Cost (CAC), and Lead-to-Customer Conversion Rate."
    )
    doc.add_paragraph(
        "A key innovation is the automated Period-over-Period Delta Engine, which splits filtered chronological data into equal halves "
        "and computes exact percentage gains or declines with directional color-coding (green for positive revenue growth, red for CAC inflation). "
        "Global multi-dimensional filters allow executives to slice metrics simultaneously by Date Range, Marketing Channel, Campaign, "
        "Customer Industry Segment, and Campaign Objective."
    )

    # Module 2
    add_custom_heading("Module 2: Customer Journey & 9-Stage Marketing Funnel", 2)
    doc.add_paragraph(
        "Rather than treating marketing and sales as disconnected silos, Module 2 models a cohesive 9-stage conversion funnel:\n"
        "1. Impressions → 2. Ad Clicks → 3. Website Visits → 4. Leads Generated → 5. Marketing-Qualified (MQL) → "
        "6. Sales-Qualified (SQL) → 7. Sales Opportunities → 8. Closed-Won Customers → 9. Attributed Revenue."
    )
    doc.add_paragraph(
        "For every stage transition, the engine calculates: (a) raw volume, (b) conversion rate from the preceding stage, (c) drop-off percentage, "
        "(d) lost prospects count, and (e) Estimated Lost Revenue Opportunity. The lost revenue metric mathematically models the pipeline value "
        "that would have materialized had the lost prospects converted at the historical downstream win probability, highlighting the exact stage "
        "inflicting the largest economic damage."
    )

    # Module 3
    add_custom_heading("Module 3: Multi-Touch Marketing Attribution Engine", 2)
    doc.add_paragraph(
        "Single-touch attribution (such as Last-Click) severely distorts marketing budget decisions by giving 100% of the credit to "
        "whichever channel was touched immediately before conversion, completely ignoring upper-funnel discovery channels. RevenueLens AI "
        "implements five transparent mathematical attribution models:"
    )
    doc.add_paragraph(
        "• First-Touch: 100% credit to the earliest discovery interaction (highlights awareness drivers like Organic & Meta Display).\n"
        "• Last-Touch: 100% credit to the final recorded marketing touchpoint prior to deal creation (favors search and direct email).\n"
        "• Linear: Uniform 1/N credit distributed equally across all recorded buyer journey touchpoints.\n"
        "• Position-Based (U-Shaped): 40% credit to the first touch, 40% to the lead conversion touch, and 20% distributed across intermediate nurture steps.\n"
        "• Time-Decay: Exponential half-life weighting: w_i = 2^(-Δt / 7 days), weighting recent touches higher while crediting earlier touches."
    )
    doc.add_paragraph(
        "Crucially, the attribution engine enforces strict normalization: the sum of weights for each customer journey is guaranteed to equal 1.0, "
        "ensuring that the total attributed revenue across channels exactly matches audited closed sales revenue."
    )

    # Module 4
    add_custom_heading("Module 4: Campaign Performance & Strategic Quadrants", 2)
    doc.add_paragraph(
        "Module 4 evaluates campaign effectiveness across Google Ads, Meta Ads, LinkedIn Ads, Email Marketing, Organic Search, Content Marketing, "
        "and Referral Programs. Beyond computing CPL and CAC, the system introduces a proprietary MBA Strategic Marketing Quadrant classification:"
    )
    doc.add_paragraph(
        "1. Vanity Trap (High Lead Volume, Low Conversion): Campaigns that produce large lead volumes at deceptive CPLs, but yield negligible customers.\n"
        "2. Hidden Gem (Low Volume, High Intent & High ROI): Campaigns with modest volume but high lead scores, quick sales cycles, and stellar ROAS.\n"
        "3. Money Pit (High Spend, Severe Negative ROI): Significant capital expenditures with inadequate pipeline contribution; primary candidates for budget pruning.\n"
        "4. Scale Candidate (Top Performer): Robust volume, high qualification rates, and positive unit economics; primary destination for expanded capital allocation."
    )

    # Module 5
    add_custom_heading("Module 5: Lead Intelligence & Machine Learning Engine", 2)
    doc.add_paragraph(
        "To understand why leads fail to convert, Module 5 trains supervised machine learning classification models (Logistic Regression "
        "and Random Forest). The predictive pipeline incorporates rigorous Data Leakage Prevention by strictly excluding all post-conversion "
        "sales outcome attributes (such as sales stage, close date, deal value, or opportunity ID) from the feature matrix."
    )
    doc.add_paragraph(
        "The model is evaluated using stratified 75/25 train/test splits, reporting Accuracy, Precision, Recall, F1-Score, Confusion Matrices, "
        "and ROC-AUC curves. Feature importance analysis isolates the primary drivers of win probability (e.g., Lead Score and Follow-Up Response Time). "
        "The module also includes an interactive real-time Lead Scorer that assigns incoming leads to High (>65%), Medium (35-65%), or Low (<35%) "
        "conversion tiers with automated SDR routing recommendations."
    )

    # Module 6
    add_custom_heading("Module 6: AI-Powered Revenue Leakage Detection & Reporting", 2)
    doc.add_paragraph(
        "Module 6 acts as an automated management consultant. A transparent heuristic rules engine scans the marketing and sales datasets "
        "for 10 distinct commercial leakage patterns, including: (1) Lead Bloat without conversion, (2) Capital Bleed / negative ROAS, "
        "(3) MQL-to-SQL qualification bottlenecks, (4) Speed-to-lead follow-up latency (>24-48 hours), (5) High revenue concentration risk, "
        "and (6) Volume vs. Value divergence."
    )
    doc.add_paragraph(
        "Each detected issue is reported with: Problem Identified, Empirical Evidence & Metrics, Likely Contributing Factors, Business Impact, "
        "Recommended Corrective Action, Priority Level (High/Med/Low), and Key KPI to Monitor. The module operates via dual synthesis: a 100% "
        "offline deterministic narrative engine, plus an optional OpenAI LLM integration. It also features a one-click ReportLab generator that "
        "compiles a C-suite ready PDF executive brief."
    )

    # Module 7
    add_custom_heading("Module 7: What-If Revenue Scenario Simulator", 2)
    doc.add_paragraph(
        "The Scenario Simulator enables executive what-if financial modeling. Leadership can manipulate strategic levers—such as scaling "
        "overall marketing budget, reallocating capital to top-ROAS channels, reducing CPL through form optimization, boosting qualification "
        "rates, accelerating sales win rates, or expanding average deal size."
    )
    doc.add_paragraph(
        "The simulation engine calculates the projected impact on Leads Generated, Closed Customers, Attributed Revenue, CAC, ROAS, and Net Profit. "
        "Calculations maintain strict mathematical consistency, preventing double-counting improvements across sequential funnel stages and "
        "presenting baseline-versus-scenario comparative ledgers."
    )

    # Module 8 & 9
    add_custom_heading("Module 8 & 9: Data Governance, Upload & Methodology", 2)
    doc.add_paragraph(
        "Module 8 provides an enterprise data management portal supporting CSV and Excel uploads for Campaigns, Leads, Customer Journey, "
        "and Sales datasets. Uploaded files undergo rigorous schema validation, type parsing, negative-value detection, and referential integrity "
        "checks. Data is persisted to SQLite, with sample benchmark resets and downloadable clean templates available."
    )
    doc.add_paragraph(
        "Module 9 establishes academic governance, documenting all mathematical formulas, attribution algorithms, machine learning validation "
        "practices, and explicit analytical disclaimers distinguishing statistical correlation from causal incrementality."
    )

    # -------------------------------------------------------------
    # 4. EMPIRICAL BENCHMARK FINDINGS
    # -------------------------------------------------------------
    add_custom_heading("4. Key Empirical Findings & Strategic Insights", 1)

    doc.add_paragraph(
        "Analysis of the benchmark multi-channel marketing dataset generated by RevenueLens AI revealed three profound commercial insights "
        "directly applicable to executive growth strategy:"
    )

    doc.add_paragraph(
        "1. The Meta Lead Gen 'Vanity Trap':\n"
        "Campaign CMP-003 ('Meta Broad Lead Gen Form') generated the single highest volume of leads in the portfolio (580 leads, representing "
        "over 28% of all leads captured) at a very attractive Cost Per Lead of $65.95. However, due to minimal qualification friction on the form, "
        "leads averaged a low lead score (38/100). Only 18 closed deals materialized, yielding an anemic conversion rate of 3.10% and an ROAS of "
        "only 0.82x (a net financial loss of $6,850). The channel created the illusion of marketing productivity while draining sales capacity."
    )

    doc.add_paragraph(
        "2. The Google Search BOFU 'Hidden Gem':\n"
        "Conversely, Campaign CMP-001 ('Google Search High Intent BOFU') exhibited a much higher nominal CPL of $111.54, generating only 260 leads. "
        "However, because these prospects were searching with high commercial purchase intent, they exhibited an average lead score of 78/100. "
        "This campaign produced 64 closed customers (24.62% conversion rate) and $114,200 in recognized revenue, generating a dominant ROAS of "
        "3.94x and an ROI of +293.8%."
    )

    doc.add_paragraph(
        "3. Speed-to-Lead Follow-Up Latency Decays Win Rates by Over 60%:\n"
        "Machine learning feature importance analysis identified SDR Follow-Up Response Time as a top predictor of deal conversion. "
        "Leads contacted within 4 hours achieved an average conversion rate of 28.4%, whereas leads contacted after 24 hours experienced a "
        "precipitous drop to 10.6% win rate—a 62.7% decay in inbound marketing ROI caused purely by sales latency."
    )

    # -------------------------------------------------------------
    # 5. PRIORITIZED ACTION ROADMAP
    # -------------------------------------------------------------
    add_custom_heading("5. Prioritized 30/60/90-Day Strategic Roadmap", 1)

    t_road = doc.add_table(rows=1, cols=4)
    t_road.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_headers = ["Horizon", "Priority", "Recommended Strategic Action", "Key Performance Indicator"]
    for i, h in enumerate(r_headers):
        cell = t_road.rows[0].cells[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    roadmap_data = [
        ("Immediate (0 - 30 Days)", "HIGH", "Reallocate 30% of marketing budget from Meta Broad Lead Gen to Google Search BOFU and Referral Programs.", "Blended Portfolio ROAS"),
        ("Immediate (0 - 30 Days)", "HIGH", "Enforce strict SDR Follow-up SLA: require outbound call/email within < 2 hours for high-intent leads.", "Median Time-to-First-Touch"),
        ("Immediate (0 - 30 Days)", "HIGH", "Add qualifying friction fields (Business Email, Company Size, Budget) to top-of-funnel web forms.", "MQL-to-SQL Acceptance Rate"),
        ("Medium-Term (30 - 60 Days)", "MEDIUM", "Deploy automated Machine Learning lead scoring in CRM to route High-Tier (>65% prob) leads to Senior AEs.", "High-Tier Opportunity Win Rate"),
        ("Medium-Term (30 - 60 Days)", "MEDIUM", "Implement structured email nurture tracks for Medium-Tier leads instead of burning SDR capacity.", "SDR Pipeline Productivity"),
        ("Strategic (60 - 90 Days)", "MEDIUM", "Transition executive board reporting from Last-Touch attribution to Position-Based (U-Shaped) attribution.", "Cross-Channel Budget Efficiency")
    ]

    for row_idx, (horiz, prio, act, kpi) in enumerate(roadmap_data):
        row = t_road.add_row()
        bg_col = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate([horiz, prio, act, kpi]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, 100, 100, 140, 140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c_idx == 1 and val == "HIGH":
                r_p = p.add_run(val)
                r_p.font.bold = True
                r_p.font.color.rgb = RGBColor(185, 28, 28)
            else:
                p.add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 6. HOW TO RUN & VERIFY
    # -------------------------------------------------------------
    add_custom_heading("6. Execution, Verification & Testing Instructions", 1)

    doc.add_paragraph(
        "The application has been verified for complete local execution, error-free navigation, and passed 100% of automated unit tests. "
        "To inspect or run the application locally, use the commands below:"
    )

    doc.add_paragraph("1. Install Dependencies:\n   pip install -r requirements.txt")
    doc.add_paragraph("2. Execute Automated Unit Tests:\n   python -m pytest -v")
    doc.add_paragraph("3. Launch the Interactive Web Dashboard:\n   streamlit run app.py")
    doc.add_paragraph("4. Access in Web Browser:\n   Navigate to http://localhost:8501")

    # Output file path
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "RevenueLens_AI_Project_Report.docx")
    doc.save(output_path)
    print(f"Report generated successfully: {output_path}")
    return output_path


if __name__ == "__main__":
    create_report()
